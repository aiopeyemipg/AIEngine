"""
make_docs.py  —  the guide as a PDF and the presentation as PowerPoint + PDF.

  GUIDE.md   -> GUIDE.pdf     (Markdown -> styled HTML -> LibreOffice PDF export)
  SLIDES.md  -> SLIDES.pptx   (python-pptx, 16:9) -> SLIDES.pdf (LibreOffice)

In both files the markers <!-- TIMING_TABLE --> and <!-- ACCURACY_TABLE -->
are replaced by tables built from results/summary.csv, so the documents always
show the latest measured numbers.

SLIDES.md format (editable by hand):
  # Slide title                 starts a new slide (the first one is the title slide;
                                its following lines become the subtitle)
  - bullet / - - sub-bullet     bullet points (left column)
  | a | b |  table rows         a table (first row = header; |---| lines ignored)
  ![caption](path/to/img.png)   an image (right column, or full width without bullets)
  > note text                   speaker notes
The same script serves ACO_Aiengine, ABC_AI_Engine, GA_AIE and PSO_AI_Engine.

Needs markdown and python-pptx (~/.venvs/aie_plots) and LibreOffice (soffice).

Usage (from the project folder):
    .../python scripts/make_docs.py [--guide] [--slides]       (default: both)
"""
import argparse
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

PROJECT = Path(__file__).resolve().parent.parent

CSS = """
@page { size: A4; margin: 18mm 16mm 18mm 16mm; }
body { font-family: 'DejaVu Sans', Arial, sans-serif; font-size: 10pt; line-height: 1.38; color: #1d1d1d; }
h1 { font-size: 20pt; color: #0b3d6e; border-bottom: 2px solid #0b3d6e; padding-bottom: 4px; }
h2 { font-size: 14.5pt; color: #0b3d6e; margin-top: 18pt; border-bottom: 1px solid #c8d3df; }
h3 { font-size: 12pt; color: #1f5f99; margin-top: 12pt; }
code { font-family: 'DejaVu Sans Mono', monospace; font-size: 8.6pt; background: #f1f3f5; }
pre { font-family: 'DejaVu Sans Mono', monospace; font-size: 7.8pt; background: #f4f6f8;
      border: 1px solid #d7dde3; padding: 6px; line-height: 1.22; }
table { border-collapse: collapse; margin: 6pt 0; font-size: 9pt; }
th { background: #0b3d6e; color: white; padding: 3px 6px; text-align: left; }
td { border: 1px solid #c8d3df; padding: 3px 6px; vertical-align: top; }
img { max-width: 100%; }
blockquote { color: #555; border-left: 3px solid #c8d3df; margin-left: 0; padding-left: 8px; }
"""


def soffice_pdf(src, out_dir, filt=None):
    """Convert a document to PDF with LibreOffice (headless)."""
    target = "pdf" if filt is None else f"pdf:{filt}"
    env = {k: v for k, v in os.environ.items() if k not in ("LD_LIBRARY_PATH", "PYTHONPATH", "PYTHONHOME")}
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless",
                        "--convert-to", target, "--outdir", str(out_dir), str(src)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=env)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def summary_tables():
    """Markdown tables from results/summary.csv (timing and accuracy)."""
    import csv
    path = PROJECT / "results" / "summary.csv"
    if not path.exists():
        return {}
    rows = list(csv.DictReader(path.open()))
    has_py = any(num(r.get("mean_ms_python")) for r in rows)
    t = ["| Algorithm | " + ("Python ms | " if has_py else "") +
         "Original C++ ms | AIE code, 1 thread ms | AIE code, parallel ms | Threads | Speed-up vs C++ | Speed-up vs 1 thread |",
         "|---|" + ("---|" if has_py else "") + "---|---|---|---|---|---|"]
    for r in rows:
        f = lambda k, d=3: (f"{num(r.get(k)):.{d}f}" if num(r.get(k)) is not None else "–")
        t.append(f"| {r['variant']} | " + (f"{f('mean_ms_python', 1)} | " if has_py else "") +
                 f"{f('mean_ms_reference_cpp')} | {f('mean_ms_aie_serial')} | {f('mean_ms_aie_parallel')} | "
                 f"{r.get('threads', '')} | {f('speedup_parallel_vs_reference', 2)}x | {f('speedup_parallel_vs_serial', 2)}x |")
    has_pyacc = any(num(r.get("python_accuracy_pct")) for r in rows)
    a = ["| Algorithm | AIE code accuracy % | AIE code mean IoU | Reference C++ accuracy % | Reference C++ mean IoU |"
         + (" Python accuracy % |" if has_pyacc else ""),
         "|---|---|---|---|---|" + ("---|" if has_pyacc else "")]
    for r in rows:
        f = lambda k, d=2: (f"{num(r.get(k)):.{d}f}" if num(r.get(k)) is not None else "–")
        a.append(f"| {r['variant']} | {f('accuracy_pct')} | {f('mean_iou_all', 4)} | {f('ref_cpp_accuracy_pct')} | "
                 f"{f('ref_cpp_mean_iou_all', 4)} |" + (f" {f('python_accuracy_pct')} |" if has_pyacc else ""))
    images = rows[0].get("images", "") if rows else ""
    note = (f"\n\nMeasured on {images} images per algorithm, Intel i5-12500 (6 cores / 12 threads), "
            "g++ -O2, idle machine; times exclude file reading. Source: results/summary.csv.")
    return {"<!-- TIMING_TABLE -->": "\n".join(t) + note, "<!-- ACCURACY_TABLE -->": "\n".join(a)}


def fill(text):
    for k, v in summary_tables().items():
        text = text.replace(k, v)
    return text


def make_guide():
    import markdown
    md = fill((PROJECT / "GUIDE.md").read_text())
    body = markdown.markdown(md, extensions=["tables", "fenced_code", "sane_lists"])
    with tempfile.TemporaryDirectory() as tmp:
        # LibreOffice's HTML import needs absolute file URLs and explicit
        # width/height (it ignores CSS max-width): 650 px = text width
        from PIL import Image

        def fit(m):
            src_img = (PROJECT / m.group(1)).resolve()
            w, h = Image.open(src_img).size
            return f'src="{src_img.as_uri()}" width="650" height="{round(650 * h / w)}"'

        body = re.sub(r'src="(?!/|http)([^"]+)"', fit, body)
        html = f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{body}</body></html>"
        src = Path(tmp) / "GUIDE.html"
        src.write_text(html)
        soffice_pdf(src, tmp, "writer_web_pdf_Export")
        shutil.copy(Path(tmp) / "GUIDE.pdf", PROJECT / "GUIDE.pdf")
    print(f"guide  -> {PROJECT / 'GUIDE.pdf'}")


def parse_slides(text):
    slides, cur = [], None
    for line in text.splitlines():
        s = line.rstrip()
        if s.startswith("# "):
            cur = {"title": s[2:].strip(), "bullets": [], "table": [], "images": [], "notes": [], "text": []}
            slides.append(cur)
        elif cur is None or not s.strip():
            continue
        elif s.lstrip().startswith("!["):
            m = re.match(r"\s*!\[(.*?)\]\((.*?)\)", s)
            cur["images"].append((m.group(1), m.group(2)))
        elif s.lstrip().startswith("|"):
            cells = [c.strip() for c in s.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                cur["table"].append(cells)
        elif s.lstrip().startswith("- "):
            level = (len(s) - len(s.lstrip())) // 2
            txt = s.lstrip()[2:]
            while txt.startswith("- "):
                level += 1
                txt = txt[2:]
            cur["bullets"].append((level, txt))
        elif s.startswith(">"):
            cur["notes"].append(s[1:].strip())
        else:
            cur["text"].append(s.strip())
    return slides


def add_runs(par, text, size, color=None, bold=False):
    """Text with **bold** and `code` spans."""
    from pptx.util import Pt
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        run = par.add_run()
        if part.startswith("**"):
            run.text, run.font.bold = part[2:-2], True
        elif part.startswith("`"):
            run.text, run.font.name = part[1:-1], "DejaVu Sans Mono"
        else:
            run.text, run.font.bold = part, bold
        run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def make_slides():
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    NAVY, GREY, WHITE = RGBColor(0x0B, 0x3D, 0x6E), RGBColor(0x55, 0x5F, 0x6A), RGBColor(0xFF, 0xFF, 0xFF)
    slides = parse_slides(fill((PROJECT / "SLIDES.md").read_text()))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    for n, sd in enumerate(slides):
        s = prs.slides.add_slide(blank)
        if n == 0:                                          # title slide
            bg = s.shapes.add_shape(1, 0, 0, W, H)
            bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
            tb = s.shapes.add_textbox(Inches(0.8), Inches(2.3), W - Inches(1.6), Inches(1.6))
            tf = tb.text_frame; tf.word_wrap = True
            add_runs(tf.paragraphs[0], sd["title"], 36, WHITE, bold=True)
            for line in sd["text"] + [b for _, b in sd["bullets"]]:
                p = tf.add_paragraph(); add_runs(p, line, 18, RGBColor(0xD6, 0xE4, 0xF0))
            continue

        bar = s.shapes.add_shape(1, 0, 0, W, Inches(1.0))
        bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.15), W - Inches(1.0), Inches(0.8))
        add_runs(tb.text_frame.paragraphs[0], sd["title"], 26, WHITE, bold=True)
        foot = s.shapes.add_textbox(W - Inches(1.2), H - Inches(0.45), Inches(1.0), Inches(0.3))
        add_runs(foot.text_frame.paragraphs[0], f"{n + 1}", 10, GREY)

        top, left = Inches(1.25), Inches(0.5)
        has_text = bool(sd["bullets"] or sd["text"] or sd["table"])
        img_w = (W * 0.46) if has_text else (W - Inches(1.0))
        txt_w = (W - Inches(1.0) - img_w - Inches(0.3)) if sd["images"] else (W - Inches(1.0))

        y = top
        if sd["bullets"] or sd["text"]:
            n_lines = len(sd["bullets"]) + len(sd["text"])
            box_h = Inches(min(5.6, 0.48 * n_lines + 0.3)) if sd["table"] else H - top - Inches(0.6)
            tb = s.shapes.add_textbox(left, y, txt_w, box_h)
            tf = tb.text_frame; tf.word_wrap = True
            first = True
            for line in sd["text"]:
                p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
                add_runs(p, line, 17, GREY)
            size = 19 if n_lines <= 6 else (17 if n_lines <= 9 else 15)
            for level, b in sd["bullets"]:
                p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
                add_runs(p, ("• " if level == 0 else "– ") + b, size - 2 * level)
                p.level = level
                p.space_after = Pt(6)
            y = y + box_h + Inches(0.1)
        if sd["table"]:
            rows, cols = len(sd["table"]), max(len(r) for r in sd["table"])
            th = min(H - y - Inches(0.5), Inches(0.42) * rows)
            tbl = s.shapes.add_table(rows, cols, left, y, txt_w, th).table
            fs = 14 if rows <= 6 else (12 if rows <= 9 else 10)
            for r, row in enumerate(sd["table"]):
                for c in range(cols):
                    cell = tbl.cell(r, c)
                    cell.text = ""
                    add_runs(cell.text_frame.paragraphs[0], row[c] if c < len(row) else "", fs,
                             WHITE if r == 0 else None, bold=(r == 0))
        if sd["images"]:
            k = len(sd["images"])
            area_x = left + txt_w + Inches(0.3) if has_text else left
            area_h = H - top - Inches(0.6)
            each_h = (area_h - Inches(0.35) * (k - 1)) / k
            for i, (cap, path) in enumerate(sd["images"]):
                p = (PROJECT / path)
                if not p.exists():
                    continue
                pic = s.shapes.add_picture(str(p), area_x, top + i * (each_h + Inches(0.35)))
                scale = min(img_w / pic.width, (each_h - Inches(0.3)) / pic.height)
                pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
                pic.left = int(area_x + (img_w - pic.width) / 2)
                if cap:
                    cb = s.shapes.add_textbox(area_x, pic.top + pic.height, img_w, Inches(0.3))
                    cp = cb.text_frame.paragraphs[0]
                    add_runs(cp, cap, 11, GREY)
        if sd["notes"]:
            s.notes_slide.notes_text_frame.text = "\n".join(sd["notes"])

    out = PROJECT / "SLIDES.pptx"
    prs.save(out)
    with tempfile.TemporaryDirectory() as tmp:
        soffice_pdf(out, tmp)
        shutil.copy(Path(tmp) / "SLIDES.pdf", PROJECT / "SLIDES.pdf")
    print(f"slides -> {out} and SLIDES.pdf ({len(slides)} slides)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--guide", action="store_true")
    ap.add_argument("--slides", action="store_true")
    a = ap.parse_args()
    both = not (a.guide or a.slides)
    if a.guide or both:
        make_guide()
    if a.slides or both:
        make_slides()


if __name__ == "__main__":
    main()
